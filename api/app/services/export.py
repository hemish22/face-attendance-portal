"""JSON/Excel export per build-spec §9."""

import io

import pandas as pd
from openpyxl.styles import PatternFill

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
AMBER = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def build_excel(result: dict, thresholds: dict) -> bytes:
    attendance_rows = []
    for p in result["present"]:
        attendance_rows.append(
            {
                "Roll No": p["roll_no"], "Name": p["name"], "Status": "Present",
                "Best Match Score": p["best_score"], "Photos Matched": p["photos_matched"],
                "Photo Files": ", ".join(ph["filename"] for ph in p["photos"]),
                "Reviewed": any(ph["status"] == "confirmed" for ph in p["photos"]),
            }
        )
    for p in result["needs_review"]:
        attendance_rows.append(
            {
                "Roll No": p["roll_no"], "Name": p["name"], "Status": "Review",
                "Best Match Score": p["best_score"], "Photos Matched": p["photos_matched"],
                "Photo Files": ", ".join(ph["filename"] for ph in p["photos"]),
                "Reviewed": False,
            }
        )
    for a in result["absent"]:
        attendance_rows.append(
            {
                "Roll No": a["roll_no"], "Name": a["name"], "Status": "Absent",
                "Best Match Score": None, "Photos Matched": 0, "Photo Files": "", "Reviewed": False,
            }
        )

    unidentified_rows = [
        {"Label": u["label"], "Appearances": u["appearances"], "Photo Files": ", ".join(ph["filename"] for ph in u["photos"])}
        for u in result["unidentified"]
    ]

    photo_index_rows = []
    for photo in result["photos"]:
        for person in photo["people"]:
            photo_index_rows.append(
                {
                    "Photo": photo["filename"], "Person": person["label"],
                    "Type": "Member" if person.get("member_id") else "Unidentified",
                    "Status": person["status"], "Match Score": None,
                }
            )

    summary = result["summary"]
    review_rate = summary["needs_review"] / summary["members"] if summary["members"] else 0
    summary_rows = [
        {"Metric": "Members", "Value": summary["members"]},
        {"Metric": "Present", "Value": summary["present"]},
        {"Metric": "Needs Review", "Value": summary["needs_review"]},
        {"Metric": "Absent", "Value": summary["absent"]},
        {"Metric": "Unidentified", "Value": summary["unidentified"]},
        {"Metric": "Photos", "Value": summary["photos"]},
        {"Metric": "Faces Detected", "Value": summary["faces_detected"]},
        {"Metric": "Faces Skipped (Low Quality)", "Value": summary["faces_skipped_low_quality"]},
        {"Metric": "Review Rate", "Value": round(review_rate, 4)},
        *[{"Metric": f"Threshold: {k}", "Value": v} for k, v in thresholds.items()],
    ]

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame(attendance_rows).to_excel(writer, sheet_name="Attendance", index=False)
        pd.DataFrame(unidentified_rows).to_excel(writer, sheet_name="Unidentified", index=False)
        pd.DataFrame(photo_index_rows).to_excel(writer, sheet_name="Photo Index", index=False)
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Summary", index=False)

        sheet = writer.sheets["Attendance"]
        status_col = 3  # "Status" is column C (1-indexed: Roll No, Name, Status)
        fill_by_status = {"Present": GREEN, "Review": AMBER, "Absent": RED}
        for row_i, row in enumerate(attendance_rows, start=2):
            fill = fill_by_status.get(row["Status"])
            if fill:
                sheet.cell(row=row_i, column=status_col).fill = fill

    return buf.getvalue()
